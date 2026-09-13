"""报告输出：xlsx / html / markdown / json。

原项目只输出一个 JSON 文本文件，而且是在热循环里「全量读 -> 重排 -> 全量写」，
论文一多就是 O(n^2)。这里改成流程跑完一次性落盘，
并且默认输出 Excel —— 选刊本来就是要在表格里横向比较、自己再筛一遍。
"""

from __future__ import annotations

import html
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from .ccf import CCF_YEAR
from .models import Journal
from .pipeline import PipelineResult, SORT_KEYS


def _conf_cycle_display(j: Journal) -> str:
    """会议固定审稿周期展示成「约3个月」，跟 letpub 的统计周期字段分列，
    避免用户把两种性质不同的数字混在一起比较。"""
    if j.metrics.conf_cycle_months is None:
        return ""
    return f"约{j.metrics.conf_cycle_months:g}个月"


# (表头, 取值函数)
COLUMNS: list[tuple[str, Any]] = [
    ("刊名", lambda j: j.name),
    ("类型", lambda j: {"journal": "期刊", "conference": "会议",
                        "book": "图书", "repository": "预印本",
                        "other": "其它"}.get(j.venue_type, j.venue_type)),
    ("CCF等级", lambda j: j.metrics.ccf_rank or ""),
    ("CCF类别", lambda j: j.metrics.ccf_category or ""),
    ("命中论文数", lambda j: j.paper_count),
    ("命中论文总被引", lambda j: j.total_citations),
    ("分区(letpub)", lambda j: f"{j.metrics.cas_quartile}区" if j.metrics.cas_quartile else ""),
    ("IF(letpub)", lambda j: j.metrics.letpub_if),
    ("两年均被引(OpenAlex)", lambda j: _round(j.metrics.impact_2yr)),
    ("CiteScore(letpub)", lambda j: j.metrics.letpub_citescore),
    ("h-index", lambda j: j.metrics.h_index),
    ("审稿周期(letpub统计)", lambda j: j.metrics.review_period or ""),
    ("会议截稿-通知周期(固定参考)", _conf_cycle_display),
    ("录用比例", lambda j: j.metrics.acceptance or ""),
    ("SCI收录", lambda j: j.metrics.sci_index or ""),
    ("学科", lambda j: j.metrics.cas_category or ""),
    ("国内刊", lambda j: "是" if j.is_chinese else ""),
    ("OA", lambda j: _yesno(j.metrics.is_oa)),
    ("DOAJ", lambda j: _yesno(j.metrics.is_in_doaj)),
    ("APC(USD)", lambda j: j.metrics.apc_usd),
    ("出版商", lambda j: j.metrics.publisher or ""),
    ("国家", lambda j: j.metrics.country_code or ""),
    ("年发文量", lambda j: j.metrics.works_count),
    ("ISSN-L", lambda j: j.issn_l or ""),
    ("letpub链接", lambda j: j.metrics.letpub_url or ""),
    ("期刊主页", lambda j: j.metrics.homepage or ""),
]

PAPER_COLUMNS: list[tuple[str, Any]] = [
    ("刊名/会议", lambda j, p: j.name),
    ("论文标题", lambda j, p: p.title),
    ("年份", lambda j, p: p.year),
    ("被引", lambda j, p: p.cited_by),
    ("作者", lambda j, p: ", ".join(p.authors[:5])),
    ("DOI", lambda j, p: p.doi or ""),
    ("链接", lambda j, p: p.url or ""),
    ("发现来源", lambda j, p: "+".join(sorted(p.providers))),
]


def _round(value: Any, digits: int = 2) -> Any:
    if isinstance(value, (int, float)):
        return round(value, digits)
    return value


def _yesno(value: Any) -> str:
    if value is True:
        return "是"
    if value is False:
        return "否"
    return ""


def _stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def write_all(result: PipelineResult, outdir: str | Path,
              basename: str | None = None,
              formats: list[str] | None = None) -> list[Path]:
    """按需输出多种格式，返回实际生成的文件列表。

    每次运行的结果单独放一个子文件夹（outdir/<basename>/），
    而不是把所有格式散在 outdir 根目录下 —— 跑得多了根目录会堆满文件，
    分文件夹之后一次运行的产出天然聚在一起，删/挪也方便。
    """
    outdir = Path(outdir)
    basename = basename or f"journals_{_stamp()}"
    run_dir = outdir / basename
    run_dir.mkdir(parents=True, exist_ok=True)
    formats = formats or ["xlsx", "html", "json", "md"]

    written: list[Path] = []
    for fmt in formats:
        path = run_dir / f"{basename}.{fmt}"
        if fmt == "xlsx":
            if write_xlsx(result, path):
                written.append(path)
        elif fmt == "html":
            write_html(result, path)
            written.append(path)
        elif fmt == "json":
            write_json(result, path)
            written.append(path)
        elif fmt in ("md", "markdown"):
            write_markdown(result, path)
            written.append(path)
    return written


# ------------------------------------------------------------------------ JSON

def write_json(result: PipelineResult, path: str | Path) -> Path:
    path = Path(path)
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "ccf_year": CCF_YEAR,
        "query": {
            "description": result.query.describe(),
            "keywords_all": result.query.keywords_all,
            "keywords_any": result.query.keywords_any,
            "scope": result.query.scope,
            "year_from": result.query.year_from,
            "year_to": result.query.year_to,
            "providers": result.query.providers,
            "max_papers": result.query.max_papers,
        },
        "sort_key": result.sort_key,
        "stats": result.stats.to_dict(),
        "journals": [j.to_dict() for j in result.journals],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                    encoding="utf-8")
    return path


# ------------------------------------------------------------------------ Excel

def write_xlsx(result: PipelineResult, path: str | Path) -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    path = Path(path)
    wb = Workbook()

    # --- sheet 1: 期刊汇总 ---
    ws = wb.active
    ws.title = "期刊汇总"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")

    ws.append(["序号"] + [c[0] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, journal in enumerate(result.journals, 1):
        ws.append([i] + [fn(journal) for _, fn in COLUMNS])

    _autosize(ws, get_column_letter, max_width=46)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    # --- sheet 2: 论文明细 ---
    ws2 = wb.create_sheet("论文明细")
    ws2.append([c[0] for c in PAPER_COLUMNS])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    for journal in result.journals:
        for paper in journal.papers:
            ws2.append([fn(journal, paper) for _, fn in PAPER_COLUMNS])
    _autosize(ws2, get_column_letter, max_width=60)
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    # --- sheet 3: 运行信息 ---
    ws3 = wb.create_sheet("运行信息")
    ws3.append(["项", "值"])
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
    rows = [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("检索式", result.query.describe()),
        ("检索源", "+".join(result.query.providers)),
        ("排序方式", f"{result.sort_key}（{SORT_KEYS.get(result.sort_key, '')}）"),
        ("CCF 目录版本", f"{CCF_YEAR} 年版（静态快照，不会跟官方改版自动更新）"),
    ]
    for k, v in result.stats.to_dict().items():
        rows.append((k, json.dumps(v, ensure_ascii=False)
                     if isinstance(v, list) else v))
    rows.append(("指标说明", "分区/IF(letpub)/审稿周期(letpub统计)/录用比例来自 letpub，仅供参考；"
                             "两年均被引来自 OpenAlex summary_stats.2yr_mean_citedness，"
                             "口径接近 JIF 但不等于官方影响因子；"
                             "CCF等级来自 CCF 推荐目录静态快照；"
                             "会议截稿-通知周期是官方固定日程的参考值，"
                             "和期刊「审稿周期(letpub统计)」是不同性质的两个数字，不能直接比较"))
    for k, v in rows:
        ws3.append([k, v])
    _autosize(ws3, get_column_letter, max_width=90)

    wb.save(path)
    return path


def _autosize(ws: Any, get_letter: Any, max_width: int = 50) -> None:
    for idx, column in enumerate(ws.columns, 1):
        width = 6
        for cell in column:
            value = cell.value
            if value is None:
                continue
            # 中文字符按两个字宽估算，否则中文列会挤在一起
            text = str(value)
            length = sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
            width = max(width, min(length + 2, max_width))
        ws.column_dimensions[get_letter(idx)].width = width


# ------------------------------------------------------------------------- HTML

_HTML_TPL = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>候选期刊 - {title}</title>
<style>
:root {{ color-scheme: light dark; }}
body {{ font-family: -apple-system, "Segoe UI", "Microsoft YaHei", sans-serif;
       margin: 0; padding: 24px; line-height: 1.55; }}
h1 {{ font-size: 20px; margin: 0 0 4px; }}
.meta {{ color: #666; font-size: 13px; margin-bottom: 16px; }}
.meta code {{ background: rgba(127,127,127,.15); padding: 1px 5px; border-radius: 3px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
th, td {{ border: 1px solid rgba(127,127,127,.35); padding: 6px 8px; text-align: left;
         vertical-align: top; }}
th {{ background: #4472c4; color: #fff; position: sticky; top: 0; cursor: pointer;
     white-space: nowrap; }}
th:hover {{ background: #365a9c; }}
tbody tr:nth-child(even) {{ background: rgba(127,127,127,.07); }}
td.num {{ text-align: right; font-variant-numeric: tabular-nums; }}
.jname {{ font-weight: 600; }}
.ccf-a {{ background: #d4edda; color: #155724; padding: 1px 6px; border-radius: 4px; font-weight: 600; }}
.ccf-b {{ background: #fff3cd; color: #856404; padding: 1px 6px; border-radius: 4px; font-weight: 600; }}
.ccf-c {{ background: #e2e3e5; color: #383d41; padding: 1px 6px; border-radius: 4px; font-weight: 600; }}
.tag {{ display: inline-block; font-size: 11px; padding: 1px 6px; border-radius: 9px;
        background: rgba(68,114,196,.18); margin-left: 4px; }}
details {{ margin-top: 4px; }}
summary {{ cursor: pointer; color: #4472c4; font-size: 12px; }}
details ol {{ margin: 6px 0 0; padding-left: 20px; font-size: 12px; color: #555; }}
.warn {{ background: #fff4e5; border-left: 4px solid #f59e19; padding: 8px 12px;
         margin-bottom: 14px; font-size: 13px; }}
footer {{ margin-top: 24px; color: #888; font-size: 12px; }}
</style>
</head>
<body>
<h1>候选投稿期刊</h1>
<div class="meta">
  检索式 <code>{query}</code> ｜ 检索源 {providers} ｜ 排序 {sort_desc}<br>
  命中 {papers} 篇论文，聚合出 {journals} 个候选载体，耗时 {elapsed}s ｜ 生成于 {now}
</div>
{warnings}
<table id="t">
<thead><tr>{head}</tr></thead>
<tbody>
{rows}
</tbody>
</table>
<footer>
分区 / IF(letpub) / 审稿周期(letpub统计) / 录用比例来自 letpub，仅供参考。<br>
「两年均被引」为 OpenAlex 的 2yr_mean_citedness，口径接近影响因子但不是官方 JIF。<br>
CCF 等级来自 CCF 推荐目录 {ccf_year} 年版静态快照，随官方改版需要手动更新，不会自动同步。<br>
「会议截稿-通知周期」是官方固定日程的参考值，跟期刊的「审稿周期(letpub统计)」是不同性质的两个数字，
不能直接比较——前者是日程表定的，后者是用户投稿耗时统计出来的。<br>
点击表头可排序。
</footer>
<script>
// 纯前端排序，省掉一个依赖
document.querySelectorAll('#t thead th').forEach(function (th, idx) {{
  th.addEventListener('click', function () {{
    var tb = document.querySelector('#t tbody');
    var rows = Array.prototype.slice.call(tb.rows);
    var asc = th.dataset.asc !== 'true';
    th.dataset.asc = asc;
    rows.sort(function (a, b) {{
      var x = a.cells[idx].dataset.v ?? a.cells[idx].innerText;
      var y = b.cells[idx].dataset.v ?? b.cells[idx].innerText;
      var nx = parseFloat(x), ny = parseFloat(y);
      var both = !isNaN(nx) && !isNaN(ny);
      var r = both ? nx - ny : String(x).localeCompare(String(y), 'zh');
      return asc ? r : -r;
    }});
    rows.forEach(function (r) {{ tb.appendChild(r); }});
  }});
}});
</script>
</body>
</html>
"""


def write_html(result: PipelineResult, path: str | Path) -> Path:
    path = Path(path)
    e = html.escape

    head = "".join(f"<th>{e(name)}</th>" for name, _ in
                   [("序号", None)] + COLUMNS)

    numeric = {"命中论文数", "命中论文总被引", "IF(letpub)", "两年均被引(OpenAlex)",
               "CiteScore(letpub)", "h-index", "APC(USD)", "年发文量", "序号"}

    rows: list[str] = []
    for i, journal in enumerate(result.journals, 1):
        cells = [f'<td class="num" data-v="{i}">{i}</td>']
        for name, fn in COLUMNS:
            value = fn(journal)
            text = "" if value is None else str(value)
            if name == "刊名":
                inner = f'<span class="jname">{e(text)}</span>'
                if journal.metrics.homepage:
                    inner = (f'<a href="{e(journal.metrics.homepage)}" '
                             f'target="_blank" rel="noopener">{inner}</a>')
                inner += _papers_details(journal, e)
                cells.append(f"<td>{inner}</td>")
            elif name == "CCF等级" and text:
                css = {"A": "ccf-a", "B": "ccf-b", "C": "ccf-c"}.get(text, "")
                cells.append(f'<td><span class="{css}">{e(text)}</span></td>')
            elif name in ("letpub链接", "期刊主页"):
                cells.append(f'<td><a href="{e(text)}" target="_blank" '
                             f'rel="noopener">链接</a></td>' if text else "<td></td>")
            elif name in numeric:
                cells.append(f'<td class="num" data-v="{e(text)}">{e(text)}</td>')
            else:
                cells.append(f"<td>{e(text)}</td>")
        rows.append("<tr>" + "".join(cells) + "</tr>")

    warn_html = ""
    if result.stats.warnings:
        items = "".join(f"<div>· {e(w)}</div>" for w in result.stats.warnings)
        warn_html = f'<div class="warn"><b>运行提示</b>{items}</div>'

    path.write_text(_HTML_TPL.format(
        title=e(result.query.describe()[:60]),
        query=e(result.query.describe()),
        providers=e("+".join(result.query.providers)),
        sort_desc=e(SORT_KEYS.get(result.sort_key, result.sort_key)),
        papers=result.stats.papers_after_merge,
        journals=result.stats.journals,
        elapsed=round(result.stats.elapsed, 1),
        now=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ccf_year=CCF_YEAR,
        warnings=warn_html,
        head=head,
        rows="\n".join(rows),
    ), encoding="utf-8")
    return path


def _papers_details(journal: Journal, e: Any) -> str:
    if not journal.papers:
        return ""
    items = []
    for paper in journal.papers[:20]:
        label = e(paper.title)
        if paper.url:
            label = f'<a href="{e(paper.url)}" target="_blank" rel="noopener">{label}</a>'
        extra = " ".join(filter(None, [
            f"{paper.year}" if paper.year else "",
            f"被引 {paper.cited_by}" if paper.cited_by is not None else "",
        ]))
        items.append(f"<li>{label} <span class='tag'>{e(extra)}</span></li>")
    more = ""
    if len(journal.papers) > 20:
        more = f"<li>… 其余 {len(journal.papers) - 20} 篇见 Excel/JSON</li>"
    return (f"<details><summary>展开 {journal.paper_count} 篇命中论文</summary>"
            f"<ol>{''.join(items)}{more}</ol></details>")


# --------------------------------------------------------------------- Markdown

def write_markdown(result: PipelineResult, path: str | Path) -> Path:
    path = Path(path)
    cols = ["#", "刊名", "类型", "CCF", "命中", "分区", "IF",
            "两年均被引", "审稿周期/会议周期", "录用比例"]
    lines = [
        "# 候选投稿期刊",
        "",
        f"- 检索式：`{result.query.describe()}`",
        f"- 检索源：{'+'.join(result.query.providers)}",
        f"- 排序：{SORT_KEYS.get(result.sort_key, result.sort_key)}",
        f"- 命中 {result.stats.papers_after_merge} 篇论文，"
        f"聚合出 {result.stats.journals} 个候选载体，耗时 {result.stats.elapsed:.1f}s",
        f"- CCF 目录版本：{CCF_YEAR} 年版（静态快照）",
        f"- 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
    ]
    if result.stats.warnings:
        lines.append("> 运行提示：" + "；".join(result.stats.warnings))
        lines.append("")

    lines.append("| " + " | ".join(cols) + " |")
    lines.append("|" + "|".join(["---"] * len(cols)) + "|")
    for i, j in enumerate(result.journals, 1):
        m = j.metrics
        # 会议显示固定周期参考值，期刊显示 letpub 统计周期，二者不混在一起比
        cycle = (_conf_cycle_display(j) + "(固定)" if j.venue_type == "conference"
                 and m.conf_cycle_months is not None
                 else (m.review_period or "-"))
        lines.append("| " + " | ".join([
            str(i),
            j.name.replace("|", "\\|"),
            {"journal": "期刊", "conference": "会议"}.get(j.venue_type, j.venue_type),
            m.ccf_rank or "-",
            str(j.paper_count),
            f"{m.cas_quartile}区" if m.cas_quartile else "-",
            str(m.letpub_if) if m.letpub_if is not None else "-",
            f"{m.impact_2yr:.2f}" if m.impact_2yr is not None else "-",
            cycle,
            m.acceptance or "-",
        ]) + " |")

    lines += ["", "## 各刊命中论文", ""]
    for j in result.journals:
        lines.append(f"### {j.name}（{j.paper_count} 篇）")
        for p in j.papers[:15]:
            bits = [f"{p.year}" if p.year else "",
                    f"被引 {p.cited_by}" if p.cited_by is not None else ""]
            suffix = "，".join([b for b in bits if b])
            link = f"[{p.title}]({p.url})" if p.url else p.title
            lines.append(f"- {link}" + (f"（{suffix}）" if suffix else ""))
        if len(j.papers) > 15:
            lines.append(f"- … 其余 {len(j.papers) - 15} 篇见 Excel/JSON")
        lines.append("")

    lines += [
        "---",
        "",
        "分区 / IF / 审稿周期(letpub统计) / 录用比例来自 letpub，仅供参考。",
        "「两年均被引」为 OpenAlex 的 `2yr_mean_citedness`，口径接近影响因子但不是官方 JIF。",
        f"CCF 等级来自 CCF 推荐目录 {CCF_YEAR} 年版静态快照，不会随官方改版自动更新。",
        "「会议截稿-通知周期(固定)」是官方固定日程的参考值，跟期刊的「审稿周期(letpub统计)」"
        "是不同性质的两个数字，前者是日程表定的，后者是用户投稿耗时统计出来的，不能直接比较。",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path
